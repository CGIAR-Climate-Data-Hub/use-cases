#!/usr/bin/env python3
"""Build the cross-use-case ingestion matrix and (optionally) write it into the CDH tracker.

Inputs (all read-only except the tracker write):
  <slug>/data/ingestion-queue.csv        one per use-case (GCF, AgWise, …) — the canonical queues
  tracker 'Priority Data' sheet          the CDH team's own list (B4T + shared assets)
  https://cgiar-climate-data-hub.github.io/catalog.json      published Hub records
  s3://digital-atlas/cdh/data/           data-lake prefixes (anonymous ListObjects)
  hub-submissions/drafts/<id>/           draft records + verify reports in this repo

Output:
  hub-submissions/matrix.csv             one row per source dataset: use-case flags, products,
                                         route, queue status, Hub status + link, draft/report links
  tracker sheet 'Use-case matrix'        the same table, with hyperlinks (--tracker PATH --write-tracker)

Rows from different sources are joined on a canonical dataset key (ALIASES below + a normalised
name). A dataset that two use-cases derive different products from is ONE row with both flags and
both products listed — that is the point of the matrix.

    python hub-submissions/build_matrix.py                       # matrix.csv only
    python hub-submissions/build_matrix.py --tracker "$XLSX" --write-tracker   # + tracker tab
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import glob
import json
import re
import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "matrix.csv"
DRAFTS = Path(__file__).resolve().parent / "drafts"
REPO_URL = "https://github.com/CGIAR-Climate-Data-Hub/use-cases/blob/main"
CATALOG_JSON = "https://cgiar-climate-data-hub.github.io/catalog.json"
CATALOG_RECORD = "https://cgiar-climate-data-hub.github.io/catalog/{id}/"
LAKE_LIST = "https://digital-atlas.s3.amazonaws.com/?list-type=2&prefix=cdh/data/&delimiter=/"
LAKE_URL = "s3://digital-atlas/cdh/data/{prefix}"

# Canonical dataset key ← every spelling seen in the queues, the tracker, the catalog and the lake.
# Keys are lowercase kebab. Add a line here when a new spelling appears; never rename a key.
ALIASES = {
    "chirps-v3-daily": ["chirps v3 - era5", "chirps v3 daily precipitation", "chirps v3.0 daily", "chirps-v3-daily", "chirps v3 daily"],
    "chirts-era5-daily": ["chirts - era5", "chirts-era5 daily", "chirts-era5-daily"],
    "chirps-v2-daily": ["chirps v2.0 daily"],
    "nex-gddp-cmip6": ["nexgddp cmip6", "nex-gddp-cmip6"],
    "agera5": ["agera5", "observed climate data - agera5"],
    "nasa-power": ["nasa power"],
    "gaez-v5": ["gaez v5"],
    "cropgrids": ["cropgrids"],
    "aqueduct-4": ["aqueduct 4.0", "wri aqueduct 4.0"],
    "mapspam-2020": ["mapspam2020", "mapspam 2020", "spam2020", "mapspam2020-v2r2"],
    "glw4-2020": ["glw4", "glw4-2020", "gridded livestock density for 2020 (glw4)"],
    "worldpop": ["worldpop population", "worldpop"],
    "esa-worldcover-2021": ["esa worldcover (admin-summary product)", "esa worldcover 2021 v200 (cropland-fraction grid)", "esa worldcover"],
    "wb-boundaries-gad": ["world bank admin boundaries", "wb-boundaries-gad"],
    "geoboundaries": ["geoboundaries (gbopen)"],
    "jrc-glofas": ["glofas flood risk", "jrc-glofas-v2.1.2"],
    "cropland-ghg-2020": ["cropland ghg emissions"],
    "soilgrids-2": ["soilgrids 2.0"],
    "enso-driver-roni": ["noaa cpc relative oceanic nino index (roni)", "enso-driver-roni"],
    "africa-precipitation-monthly-seasonal": ["monthly and seasonal precipitation from chirps v3 (africa)", "africa-precipitation-monthly-seasonal"],
    "atlas-hazards": ["adaptation atlas hazards"],
}
# Use-cases that use a Priority-Data dataset without a queue row of their own (the GCF review's
# "already held" list; see gcf-preparation-facility/data/ingestion-queue.csv note block).
EXTRA_USES = {
    "GCF": ["aqueduct-4", "mapspam-2020", "glw4-2020", "worldpop", "wb-boundaries-gad", "cropland-ghg-2020",
            "nex-gddp-cmip6", "chirps-v3-daily", "chirts-era5-daily", "agera5"],
    # AgWise: agwise-data reads staged AgERA5 / CHIRPS v3; brief + aggeodata use CHIRTS, NASA POWER;
    # WB boundaries = the Hub standard AgWise's geoBoundaries row must align with (agwise/BRIEF.md).
    "AgWise": ["agera5", "chirps-v3-daily", "chirts-era5-daily", "nasa-power", "wb-boundaries-gad"],
}
STATUS_RANK = {"published": 0, "submitted": 1, "verified": 2, "draft": 3, "in data lake, no record": 4, "in-progress": 5,
               "todo": 6, "queued": 6, "deferred": 7}


def norm(name: str) -> str:
    s = re.sub(r"\s+", " ", str(name or "")).strip().lower()
    return s


def key_for(name: str) -> str:
    n = norm(name)
    for k, al in ALIASES.items():
        if n == k or n in al:
            return k
    # fallback: kebab of the name without parenthetical qualifiers
    base = re.sub(r"\(.*?\)", "", n)
    base = re.sub(r"[^a-z0-9]+", "-", base).strip("-")
    return base or n


def load_queues() -> list[dict]:
    rows = []
    for p in sorted(glob.glob(str(ROOT / "*" / "data" / "ingestion-queue.csv"))):
        slug = Path(p).parents[1].name
        uc = {"gcf-preparation-facility": "GCF", "agwise": "AgWise", "b4t": "B4T"}.get(slug, slug)
        with open(p) as f:
            for r in csv.DictReader(f):
                r["_use_case"] = uc
                r["_src"] = str(Path(p).relative_to(ROOT))
                rows.append(r)
    return rows


def load_priority(tracker: Path | None) -> list[dict]:
    if not tracker or not tracker.exists():
        return []
    import openpyxl  # noqa: PLC0415
    ws = openpyxl.load_workbook(tracker, read_only=True)["Priority Data"]
    hdr = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        d = dict(zip(hdr, r))
        d["_use_case"] = "Priority Data"
        out.append(d)
    return out


def load_catalog() -> dict[str, dict]:
    try:
        d = requests.get(CATALOG_JSON, timeout=40).json()
    except Exception as e:  # noqa: BLE001
        print(f"warn: catalog.json unavailable ({e.__class__.__name__})", file=sys.stderr)
        return {}
    out = {}
    for x in d.get("dataset", []):
        ids = x.get("identifier")
        ids = ids if isinstance(ids, list) else [ids]
        rid = next((i for i in ids if i and not str(i).startswith("http")), None)
        if rid:
            out[key_for(rid)] = {"id": rid, "name": x.get("name"), "url": x.get("url") or CATALOG_RECORD.format(id=rid)}
            out.setdefault(key_for(x.get("name", "")), out[key_for(rid)])
    return out


def load_lake() -> dict[str, str]:
    try:
        r = requests.get(LAKE_LIST, timeout=40)
        root = ET.fromstring(r.text)
        ns = root.tag.split("}")[0] + "}" if "}" in root.tag else ""
        prefixes = [p.find(f"{ns}Prefix").text for p in root.iter(f"{ns}CommonPrefixes")]
    except Exception as e:  # noqa: BLE001
        print(f"warn: lake listing unavailable ({e.__class__.__name__})", file=sys.stderr)
        return {}
    out = {}
    for p in prefixes:
        name = p.rstrip("/").split("/")[-1]
        out[key_for(name)] = LAKE_URL.format(prefix=name + "/")
    return out


def load_drafts() -> dict[str, dict]:
    out = {}
    for d in sorted(DRAFTS.glob("*/")):
        rid = d.name
        y = d / f"{rid}.yaml"
        if not y.exists():
            continue
        rep = d / "verify-report.md"
        status = "draft"
        issue = None
        if rep.exists():
            txt = rep.read_text()
            status = "verified" if re.search(r"\*\*HIGH 0\*\*|HIGH 0 ", txt) and "HIGH" in txt else "draft"
        meta = d / "SUBMISSION.md"
        if meta.exists():
            m = re.search(r"issue:\s*(\S+)", meta.read_text())
            if m:
                issue, status = m.group(1), "submitted"
        out[key_for(rid)] = {"id": rid, "yaml": f"{REPO_URL}/hub-submissions/drafts/{rid}/{rid}.yaml",
                             "report": f"{REPO_URL}/hub-submissions/drafts/{rid}/verify-report.md" if rep.exists() else "",
                             "status": status, "issue": issue or ""}
    return out


def build(tracker: Path | None) -> list[dict]:
    queues = load_queues()
    prio = load_priority(tracker)
    catalog, lake, drafts = load_catalog(), load_lake(), load_drafts()
    rows: dict[str, dict] = {}

    def row(key):
        return rows.setdefault(key, {"key": key, "dataset": "", "uses": set(), "products": [], "routes": set(), "storage": set(),
                                     "queue_status": [], "priority": [], "licence": set(), "leads": set(), "locations": [],
                                     "notes": [], "sources": []})

    for r in queues:
        k = key_for(r["Dataset"])
        x = row(k)
        x["dataset"] = x["dataset"] or re.sub(r"\s*\((admin-summary product|cropland-fraction grid)\)", "", r["Dataset"])
        x["uses"].add(r["_use_case"])
        if "(" in r["Dataset"] and "product" in r["Dataset"] or "grid" in r["Dataset"]:
            x["products"].append(f"{r['_use_case']}: {r['Dataset']}")
        x["routes"].add(r.get("Route", "")); x["storage"].add(r.get("Storage", ""))
        x["queue_status"].append(f"{r['_use_case']}:{r.get('Status', '')}")
        x["priority"].append(f"{r['_use_case']}:{r.get('Review priority') or r.get('Priority') or ''}")
        if r.get("Licence"):
            x["licence"].add(r["Licence"])
        if r.get("Proposed lead"):
            x["leads"].add(r["Proposed lead"])
        if r.get("Location"):
            x["locations"].append(r["Location"])
        x["sources"].append(r["_src"])
    for r in prio:
        k = key_for(r["Dataset"])
        x = row(k)
        x["dataset"] = x["dataset"] or str(r["Dataset"])
        for u in str(r.get("Use Case") or "").replace("+", ",").split(","):
            u = u.strip()
            if u and u.lower() not in ("more", "none"):
                x["uses"].add({"b4t": "B4T", "gcf": "GCF", "agwise": "AgWise"}.get(u.lower(), u))
        x["uses"].add("Priority Data")
        if r.get("Storage"):
            x["storage"].add(str(r["Storage"]))
        x["queue_status"].append(f"Priority Data:{r.get('Status', '')}")
        if r.get("Location"):
            x["locations"].append(str(r["Location"]))
        if r.get("Uploader/Person"):
            x["leads"].add(str(r["Uploader/Person"]))
        x["sources"].append("tracker:Priority Data")
    for uc, keys in EXTRA_USES.items():
        for k in keys:
            if k in rows:
                rows[k]["uses"].add(uc)

    out = []
    for k, x in rows.items():
        cat, lk, dr = catalog.get(k), lake.get(k), drafts.get(k)
        if cat:
            hub_status, hub_url = "published", cat["url"]
            x["dataset"] = cat["name"]  # the catalog's title is the canonical spelling once published
        elif dr:
            hub_status, hub_url = dr["status"], ""
        elif lk:
            hub_status, hub_url = "in data lake, no record", ""
        else:
            st = [s.split(":", 1)[1] for s in x["queue_status"]]
            hub_status = "in-progress" if "in-progress" in st else ("queued" if any(s == "todo" for s in st) else (st[0] if st else "queued"))
            hub_url = ""
        uses = sorted(x["uses"] - {"Priority Data"}, key=str.lower)
        out.append({
            "dataset": x["dataset"], "key": k,
            "GCF": "x" if "GCF" in x["uses"] else "", "AgWise": "x" if "AgWise" in x["uses"] else "",
            "B4T": "x" if "B4T" in x["uses"] else "", "other_use_cases": "; ".join(u for u in uses if u not in ("GCF", "AgWise", "B4T")),
            "n_use_cases": len(uses), "products": "; ".join(x["products"]),
            "route": "/".join(sorted(r for r in x["routes"] if r)), "storage": "/".join(sorted(s for s in x["storage"] if s)),
            "hub_status": hub_status, "hub_record_url": hub_url, "lake_prefix": lk or "",
            "draft_record": dr["yaml"] if dr else "", "verify_report": dr["report"] if dr else "", "submission_issue": dr["issue"] if dr else "",
            "queue_status": "; ".join(x["queue_status"]), "priority": "; ".join(p for p in x["priority"] if not p.endswith(":")),
            "licence": " | ".join(sorted(x["licence"]))[:160], "proposed_lead": "; ".join(sorted(x["leads"])),
            "location": " ; ".join(dict.fromkeys(x["locations"]))[:300], "sources": "; ".join(dict.fromkeys(x["sources"])),
        })
    out.sort(key=lambda r: (STATUS_RANK.get(r["hub_status"], 9), -r["n_use_cases"], r["dataset"].lower()))
    return out


COLS = ["dataset", "key", "GCF", "AgWise", "B4T", "other_use_cases", "n_use_cases", "products", "route", "storage", "hub_status",
        "hub_record_url", "lake_prefix", "draft_record", "verify_report", "submission_issue", "queue_status", "priority", "licence",
        "proposed_lead", "location", "sources"]


def write_csv(rows: list[dict]):
    with open(OUT, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print(f"matrix.csv: {len(rows)} rows → {OUT}")


def write_tracker(rows: list[dict], tracker: Path, replace_tabs: tuple[str, ...] = ("GCF use-case", "AgWise use-case")):
    import openpyxl  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font  # noqa: PLC0415
    bak = tracker.with_name(f"{tracker.stem}.bak-{dt.date.today().isoformat()}-matrix{tracker.suffix}")
    shutil.copy2(tracker, bak)
    wb = openpyxl.load_workbook(tracker)
    before = {n: [[c.value for c in r] for r in wb[n].iter_rows()] for n in wb.sheetnames if n not in replace_tabs and n != "Use-case matrix"}
    for t in replace_tabs + ("Use-case matrix",):
        if t in wb.sheetnames:
            del wb[t]
    ws = wb.create_sheet("Use-case matrix", index=1)
    ws.append(COLS)
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r[c] for c in COLS])
        rr = ws.max_row
        for col in ("hub_record_url", "draft_record", "verify_report"):
            v = r[col]
            if v:
                cell = ws.cell(rr, COLS.index(col) + 1); cell.hyperlink = v; cell.style = "Hyperlink"
    ws.append([])
    ws.append([f"Generated {dt.datetime.now():%Y-%m-%d %H:%M} by hub-submissions/build_matrix.py from the use-case queues "
               f"(<slug>/data/ingestion-queue.csv), this workbook's Priority Data, catalog.json and the data lake. Canonical copy: "
               f"{REPO_URL}/hub-submissions/matrix.csv. Do not hand-edit; edit the use-case queue and regenerate."])
    widths = {"A": 40, "B": 26, "C": 6, "D": 8, "E": 6, "F": 14, "G": 6, "H": 44, "I": 18, "J": 12, "K": 22, "L": 40, "M": 34, "N": 40,
              "O": 40, "P": 14, "Q": 36, "R": 22, "S": 34, "T": 22, "U": 50, "V": 40}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "B2"
    wb.save(tracker)
    wb2 = openpyxl.load_workbook(tracker)
    after = {n: [[c.value for c in r] for r in wb2[n].iter_rows()] for n in before}
    assert before == after, "a sheet other than the matrix/use-case tabs changed — restore from backup"
    print(f"tracker: wrote 'Use-case matrix' ({len(rows)} rows), removed {replace_tabs}; backup {bak.name}; other sheets unchanged")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--tracker", help="path to CDH_data-ingestion v2.xlsx (read Priority Data; written with --write-tracker)")
    ap.add_argument("--write-tracker", action="store_true")
    a = ap.parse_args(argv)
    tracker = Path(a.tracker) if a.tracker else None
    rows = build(tracker)
    write_csv(rows)
    if a.write_tracker and tracker:
        write_tracker(rows, tracker)


if __name__ == "__main__":
    main()
